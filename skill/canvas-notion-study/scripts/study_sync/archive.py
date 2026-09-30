"""Small, recoverable file archive used by the Canvas collector.

The archive deliberately has no Canvas-specific network code.  A caller gives
this module a client with a ``get_bytes`` method (the bundled
``CanvasClient`` does) and receives a JSON-serialisable file record back.  All
writes go through a temporary file and an atomic replace, so an interrupted
download cannot destroy the last good copy.
"""

from __future__ import annotations

import hashlib
import io
import json
import os
from pathlib import Path
import re
import tempfile
from typing import Any, Callable, Mapping
import unicodedata
import zipfile


DEFAULT_MAX_FILE_SIZE = 50 * 1024 * 1024


class ArchiveError(RuntimeError):
    """Base class for archive failures."""


class FileTooLargeError(ArchiveError):
    """Raised when a download exceeds the configured byte limit."""


def safe_component(value: Any, *, fallback: str = "item", max_length: int = 96) -> str:
    """Return a filesystem-safe, bounded path component.

    This function intentionally does not preserve path separators.  Canvas
    display names are user-controlled and may contain ``..``, slashes, or
    Unicode control characters.
    """

    text = unicodedata.normalize("NFKC", str(value or "")).strip()
    text = re.sub(r"[\x00-\x1f\x7f]", "_", text)
    text = re.sub(r"[^A-Za-z0-9._-]+", "_", text)
    text = text.strip("._")
    if not text or text in {".", ".."}:
        text = fallback
    return text[:max_length] or fallback


def source_id_for_file(file_record: Mapping[str, Any]) -> str:
    """Get a stable source ID for a Canvas file or an HTML file reference."""

    explicit = file_record.get("source_id")
    if explicit:
        return str(explicit)
    file_id = file_record.get("id")
    if file_id is not None and str(file_id).strip():
        return f"file:{file_id}"
    source_url = file_record.get("download_url") or file_record.get("url") or file_record.get("source_url")
    if source_url:
        digest = hashlib.sha256(str(source_url).encode("utf-8", "replace")).hexdigest()[:20]
        return f"file-url:{digest}"
    name = file_record.get("display_name") or file_record.get("filename") or file_record.get("name") or "file"
    digest = hashlib.sha256(str(name).encode("utf-8", "replace")).hexdigest()[:20]
    return f"file-name:{digest}"


def _file_content_type(file_record: Mapping[str, Any]) -> str | None:
    for key, value in file_record.items():
        if str(key).lower() in {"content_type", "content-type", "mime_type", "mime"} and value:
            return str(value)
    return None


def archive_file_path(
    archive_dir: os.PathLike[str] | str,
    source_id: str,
    filename: str,
    *,
    version: str | None = None,
) -> Path:
    """Return the deterministic destination path for one source file.

    A short digest protects against collisions after sanitising a source ID;
    the readable prefix still makes the archive easy to browse by hand.
    """

    root = Path(archive_dir).expanduser()
    sid = str(source_id)
    sid_digest = hashlib.sha256(sid.encode("utf-8", "replace")).hexdigest()[:12]
    sid_part = safe_component(sid, fallback="source", max_length=72)
    name = safe_component(filename, fallback="file", max_length=180)
    source_dir = root / "files" / f"{sid_part}-{sid_digest}"
    if version:
        source_dir = source_dir / safe_component(version, fallback="version", max_length=40)
    return source_dir / name


def _previous_meta(previous: Any, source_id: str) -> Mapping[str, Any] | None:
    """Find a prior file record in common snapshot/record shapes."""

    if previous is None:
        return None
    if isinstance(previous, Mapping):
        if source_id_for_file(previous) == source_id and (
            previous.get("local_path") or previous.get("download_status")
        ):
            return previous
        for key in ("files", "records"):
            value = previous.get(key)
            found = _previous_meta(value, source_id)
            if found:
                return found
        for value in previous.values():
            found = _previous_meta(value, source_id)
            if found:
                return found
        return None
    if isinstance(previous, (list, tuple)):
        for value in previous:
            found = _previous_meta(value, source_id)
            if found:
                return found
    return None


def _source_marker(file_record: Mapping[str, Any]) -> tuple[Any, ...]:
    """Fields that identify a Canvas version when deciding whether to reuse."""

    return (
        file_record.get("updated_at") or file_record.get("modified_at"),
        file_record.get("size"),
        file_record.get("md5"),
    )


def _version_token(file_record: Mapping[str, Any], data: bytes | None = None) -> str | None:
    """Derive a stable immutable version directory from Canvas metadata/content."""

    marker = _source_marker(file_record)
    if any(value not in (None, "") for value in marker):
        encoded = json.dumps(marker, ensure_ascii=False, sort_keys=True, default=str).encode("utf-8")
        return hashlib.sha256(encoded).hexdigest()[:20]
    if data is not None:
        return hashlib.sha256(data).hexdigest()[:20]
    return None


def _existing_copy_is_reusable(destination: Path, file_record: Mapping[str, Any], previous: Mapping[str, Any] | None) -> bool:
    if not destination.is_file():
        return False
    if previous is None:
        # A second reference in the same collection run has the same
        # deterministic destination.  Reusing it prevents duplicate writes.
        return True
    prior_path = previous.get("local_path")
    try:
        if prior_path and Path(str(prior_path)).resolve() == destination.resolve():
            old_marker = previous.get("source_marker")
            current_marker = _source_marker(file_record)
            return old_marker is None or tuple(old_marker) == current_marker
    except (OSError, RuntimeError, TypeError):
        return False
    return False


def _prior_destination(
    archive_dir: os.PathLike[str] | str,
    source_id: str,
    filename: str,
    previous: Mapping[str, Any] | None,
    *,
    version: str | None = None,
) -> Path | None:
    """Reuse a prior path only inside the current immutable version directory.

    Older snapshots may contain an unversioned path.  That path is eligible
    only for another unversioned lookup; a newly observed Canvas version must
    always get a new directory so a stable Notion link to the old bytes keeps
    working.
    """

    if not previous or not previous.get("local_path"):
        return None
    current = archive_file_path(archive_dir, source_id, filename, version=version)
    try:
        candidate = Path(str(previous["local_path"])).expanduser()
        if candidate.parent.resolve() == current.parent.resolve():
            return candidate
    except (OSError, RuntimeError, TypeError):
        pass
    return None


def _destination_version_token(
    archive_dir: os.PathLike[str] | str,
    source_id: str,
    filename: str,
    destination: Path,
) -> str | None:
    """Return the token encoded by a versioned destination, if any."""

    source_dir = archive_file_path(archive_dir, source_id, filename).parent.resolve()
    try:
        if destination.parent.parent.resolve() == source_dir:
            return destination.parent.name
    except (OSError, RuntimeError):
        return None
    return None


def _existing_versioned_destination(
    archive_dir: os.PathLike[str] | str,
    source_id: str,
    filename: str,
) -> Path | None:
    """Find a prior content-hash version when no Canvas metadata is present."""

    base = archive_file_path(archive_dir, source_id, filename).parent
    try:
        candidates = [path for path in base.glob(f"*/{safe_component(filename, fallback='file')}") if path.is_file()]
    except OSError:
        return None
    return candidates[0] if len(candidates) == 1 else None


def _normalise_content_type(content_type: Any) -> str:
    return str(content_type or "").lower().split(";", 1)[0].strip()


def _document_kind(data: bytes, filename: str = "", content_type: Any = None) -> str | None:
    """Identify supported document types from metadata, extension, or magic bytes."""

    suffix = Path(filename).suffix.lower()
    mime = _normalise_content_type(content_type)
    if mime in {"application/pdf", "application/x-pdf", "pdf"}:
        return "pdf"
    if mime in {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/msword",
        "docx",
    }:
        return "docx"
    if mime in {
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "application/vnd.ms-powerpoint",
        "pptx",
    }:
        return "pptx"
    if suffix == ".pdf":
        return "pdf"
    if suffix == ".docx":
        return "docx"
    if suffix == ".pptx":
        return "pptx"
    # HTML-discovered links often have a Canvas-generated name such as
    # ``file-44339276``.  Content signatures keep those records extractable
    # even when the file metadata is unavailable.
    if data.startswith(b"%PDF-"):
        return "pdf"
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                names = set(archive.namelist())
            if any(name.startswith("word/") for name in names):
                return "docx"
            if any(name.startswith("ppt/") for name in names):
                return "pptx"
        except (OSError, zipfile.BadZipFile):
            pass
    return None


def _filename_with_document_extension(filename: str, data: bytes, content_type: Any = None) -> str:
    """Add a safe, useful extension to extensionless/generic Canvas names."""

    kind = _document_kind(data, filename, content_type)
    if not kind:
        return filename
    extension = f".{kind}"
    suffix = Path(filename).suffix.lower()
    if suffix in {"", ".bin", ".dat", ".download"}:
        stem = filename[: -len(suffix)] if suffix else filename
        return f"{stem}{extension}"
    return filename


def _document_signature_matches(data: bytes, filename: str, content_type: Any = None) -> bool:
    """Reject a prior JSON/error body masquerading as a downloaded document."""

    expected = _document_kind(b"", filename, content_type)
    if expected == "pdf":
        return data.startswith(b"%PDF-")
    if expected in {"docx", "pptx"}:
        if not data.startswith(b"PK\x03\x04"):
            return False
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                names = set(archive.namelist())
            return any(name.startswith("word/" if expected == "docx" else "ppt/") for name in names)
        except (OSError, zipfile.BadZipFile):
            return False
    return True


def _extract_pdf(data: bytes) -> tuple[str | None, str]:
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError:
        return None, "unavailable"
    try:
        reader = PdfReader(io.BytesIO(data))
        chunks = [(page.extract_text() or "").strip() for page in reader.pages]
        return "\n\n".join(chunk for chunk in chunks if chunk), "extracted"
    except Exception:
        return None, "failed"


def _extract_docx(data: bytes) -> tuple[str | None, str]:
    try:
        from docx import Document  # type: ignore
    except ImportError:
        return None, "unavailable"
    try:
        document = Document(io.BytesIO(data))
        chunks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    chunks.append(" | ".join(cells))
        return "\n".join(chunks), "extracted"
    except Exception:
        return None, "failed"


def _extract_pptx(data: bytes) -> tuple[str | None, str]:
    try:
        from pptx import Presentation  # type: ignore
    except ImportError:
        return None, "unavailable"
    try:
        presentation = Presentation(io.BytesIO(data))
        chunks: list[str] = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                text = getattr(shape, "text", "")
                if text and text.strip():
                    chunks.append(text.strip())
        return "\n".join(chunks), "extracted"
    except Exception:
        return None, "failed"


def extract_document_text(
    data: bytes,
    filename: str,
    content_type: str | None = None,
    *,
    enabled: bool = True,
) -> tuple[str | None, str]:
    """Extract text from an optional PDF, DOCX, or PPTX dependency.

    Missing optional dependencies are reported as ``unavailable`` and do not
    make collection fail.  Unsupported file types are ``unsupported``.
    """

    if not enabled:
        return None, "not_requested"
    kind = _document_kind(data, filename, content_type)
    if kind == "pdf":
        return _extract_pdf(data)
    if kind == "docx":
        return _extract_docx(data)
    if kind == "pptx":
        return _extract_pptx(data)
    return None, "unsupported"


def _hash_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def hash_file(path: os.PathLike[str] | str) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def extract_document_path(path: Path, filename: str, content_type: str | None, *, enabled=True):
    """Extract from a file path without first loading the whole download."""
    if not enabled:
        return None, "not_requested"
    with path.open("rb") as stream:
        prefix = stream.read(4096)
    kind = _document_kind(prefix, filename, content_type)
    try:
        if kind == "pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            return "\n".join(page.extract_text() or "" for page in reader.pages).strip(), "extracted"
        if kind == "docx":
            from docx import Document
            document = Document(str(path))
            parts = [paragraph.text for paragraph in document.paragraphs]
            parts.extend(" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
            return "\n".join(parts).strip(), "extracted"
        if kind == "pptx":
            from pptx import Presentation
            document = Presentation(str(path))
            return "\n".join(shape.text for slide in document.slides for shape in slide.shapes if hasattr(shape, "text")).strip(), "extracted"
        return None, "unsupported"
    except ImportError:
        return None, "unavailable"
    except Exception:
        return None, "failed"


def _archive_stream(client, url, file_record, archive_dir, filename, content_type, prior,
                    metadata_version, max_file_size, extract_documents):
    root = Path(archive_dir)
    root.mkdir(parents=True, exist_ok=True)
    fd, temporary_name = tempfile.mkstemp(prefix=".download-", dir=root)
    os.close(fd)
    temporary = Path(temporary_name)
    try:
        info = client.download_to(url, temporary, max_bytes=max_file_size, allow_external=True)
        size = temporary.stat().st_size
        if max_file_size is not None and size > max_file_size:
            raise FileTooLargeError("file exceeds configured size limit")
        digest = info.get("sha256") or hash_file(temporary)
        with temporary.open("rb") as stream:
            prefix = stream.read(4096)
        filename = _filename_with_document_extension(filename, prefix, content_type)
        if Path(filename).suffix.lower() not in {".docx", ".pptx"} and prefix.startswith(b"PK"):
            try:
                with zipfile.ZipFile(temporary) as package:
                    names = package.namelist()
                    if "word/document.xml" in names:
                        filename += ".docx"
                    elif "ppt/presentation.xml" in names:
                        filename += ".pptx"
            except zipfile.BadZipFile:
                pass
        source_id = source_id_for_file(file_record)
        version = metadata_version or digest[:20]
        target = archive_file_path(archive_dir, source_id, filename, version=version)
        status = "downloaded"
        if prior and prior.get("sha256") == digest and prior.get("local_path"):
            previous_path = Path(str(prior["local_path"]))
            if previous_path.is_file() and hash_file(previous_path) == digest:
                target, status = previous_path, "reused"
        if status != "reused":
            target.parent.mkdir(parents=True, exist_ok=True)
            os.replace(temporary, target)
        text, extraction_status = extract_document_path(target, filename, content_type, enabled=extract_documents)
        result = dict(file_record)
        result.update({"source_id": source_id, "filename": filename, "local_path": str(target),
                       "sha256": digest, "size": size, "download_status": status,
                       "extraction_status": extraction_status, "version_token": version,
                       "source_marker": list(_source_marker(file_record))})
        if content_type:
            result["content_type"] = content_type
        if text is not None:
            result["extracted_text"] = text
        result.pop("warning", None)
        return result
    finally:
        temporary.unlink(missing_ok=True)


def archive_bytes(
    data: bytes,
    archive_dir: os.PathLike[str] | str,
    source_id: str,
    filename: str,
    *,
    content_type: str | None = None,
    max_file_size: int | None = DEFAULT_MAX_FILE_SIZE,
    extract_documents: bool = True,
    destination: os.PathLike[str] | str | None = None,
    source_marker: tuple[Any, ...] | None = None,
    version: str | None = None,
) -> dict[str, Any]:
    """Atomically archive bytes and return a normalized Canvas file record."""

    if max_file_size is not None and len(data) > max_file_size:
        return {
            "source_id": str(source_id),
            "download_status": "too_large",
            "size": len(data),
            "warning": "file exceeds configured size limit",
        }

    filename = safe_component(filename, fallback="file")
    filename = _filename_with_document_extension(filename, data, content_type)
    target = Path(destination) if destination is not None else archive_file_path(archive_dir, source_id, filename, version=version)
    target = target.expanduser()
    try:
        target.parent.mkdir(parents=True, exist_ok=True)
        fd, temporary_name = tempfile.mkstemp(prefix=".download-", dir=str(target.parent))
        try:
            with os.fdopen(fd, "wb") as temporary:
                temporary.write(data)
                temporary.flush()
                os.fsync(temporary.fileno())
            os.replace(temporary_name, target)
        finally:
            try:
                os.unlink(temporary_name)
            except FileNotFoundError:
                pass
    except OSError:
        return {
            "source_id": str(source_id),
            "download_status": "failed",
            "warning": "archive write failed",
        }

    extracted_text, extraction_status = extract_document_text(
        data,
        filename,
        content_type,
        enabled=extract_documents,
    )
    record: dict[str, Any] = {
        "source_id": str(source_id),
        "filename": str(filename),
        "local_path": str(target),
        "sha256": _hash_bytes(data),
        "size": len(data),
        "download_status": "downloaded",
        "extraction_status": extraction_status,
    }
    if source_marker is not None:
        record["source_marker"] = list(source_marker)
    if version:
        record["version_token"] = str(version)
    if content_type:
        record["content_type"] = content_type
    if extracted_text is not None:
        record["extracted_text"] = extracted_text
    return record


def _download_bytes(client: Any, url: str, max_file_size: int | None) -> bytes:
    getter = getattr(client, "get_bytes", None) or getattr(client, "download", None)
    if getter is None:
        raise ArchiveError("Canvas client has no byte download method")
    try:
        try:
            value = getter(url, max_bytes=max_file_size, allow_external=True)
        except TypeError:
            try:
                value = getter(url, max_bytes=max_file_size)
            except TypeError:
                # Some clients expose download(url) without optional arguments.
                value = getter(url)
    except Exception as exc:
        # Avoid importing CanvasClient here (it imports this module).  The
        # class-name check bridges its bounded-response exception cleanly.
        if type(exc).__name__ in {"CanvasDownloadTooLarge", "FileTooLargeError"}:
            raise FileTooLargeError("file exceeds configured size limit") from exc
        raise
    if isinstance(value, tuple) and len(value) == 2:
        value = value[0]
    if not isinstance(value, (bytes, bytearray, memoryview)):
        raise ArchiveError("Canvas byte download returned an unsupported value")
    data = bytes(value)
    if max_file_size is not None and len(data) > max_file_size:
        raise FileTooLargeError("file exceeds configured size limit")
    return data


def _too_large_result(
    file_record: Mapping[str, Any],
    source_id: str,
    filename: str,
    prior: Mapping[str, Any] | None,
) -> dict[str, Any]:
    """Return a bounded per-file result, retaining a verified prior copy."""

    result = dict(file_record)
    result.update(
        {
            "source_id": source_id,
            "filename": filename,
            "download_status": "too_large",
            "warning": "file exceeds configured size limit",
        }
    )
    prior_path = prior.get("local_path") if prior else None
    if prior_path and Path(str(prior_path)).is_file():
        result.update(
            {
                "local_path": str(prior_path),
                "sha256": prior.get("sha256"),
                "warning": "file exceeds configured size limit; prior file preserved",
            }
        )
        if prior.get("version_token"):
            result["version_token"] = prior["version_token"]
        if prior.get("extracted_text") is not None:
            result["extracted_text"] = prior["extracted_text"]
            result["extraction_status"] = prior.get("extraction_status", "extracted")
    return result


def download_and_archive(
    client: Any,
    file_record: Mapping[str, Any],
    archive_dir: os.PathLike[str] | str,
    *,
    previous: Any = None,
    max_file_size: int | None = DEFAULT_MAX_FILE_SIZE,
    extract_documents: bool = True,
    downloader: Callable[[str, int | None], bytes] | None = None,
) -> dict[str, Any]:
    """Download one Canvas file and archive it safely.

    ``previous`` may be a previous file record or a full snapshot.  A failed
    refresh keeps a prior local file and marks it ``preserved``.
    """

    source_id = source_id_for_file(file_record)
    content_type = _file_content_type(file_record)
    filename = (
        file_record.get("display_name")
        or file_record.get("filename")
        or file_record.get("name")
        or f"{source_id}.bin"
    )
    filename = safe_component(filename, fallback="file")
    filename = _filename_with_document_extension(filename, b"", content_type)
    prior = _previous_meta(previous, source_id)
    metadata_version = _version_token(file_record)
    destination = archive_file_path(archive_dir, source_id, filename, version=metadata_version)
    prior_destination = _prior_destination(archive_dir, source_id, filename, prior, version=metadata_version)
    if prior_destination is not None and metadata_version is not None:
        destination = prior_destination
    existing_version: Path | None = None
    if metadata_version is None and prior is None:
        existing_version = _existing_versioned_destination(archive_dir, source_id, filename)
        if existing_version is not None:
            destination = existing_version

    # Without Canvas version metadata we must fetch once and derive a content
    # hash; otherwise a changed file could overwrite/reuse an old path.
    known_version = metadata_version is not None or existing_version is not None
    can_reuse = known_version and _existing_copy_is_reusable(destination, file_record, prior)
    if can_reuse and destination.is_file() and content_type:
        try:
            expected = _document_kind(b"", filename, content_type)
            if expected in {"docx", "pptx"}:
                with zipfile.ZipFile(destination) as package:
                    can_reuse = any(name.startswith("word/" if expected == "docx" else "ppt/") for name in package.namelist())
            else:
                with destination.open("rb") as existing:
                    can_reuse = _document_signature_matches(existing.read(4096), filename, content_type)
        except (OSError, zipfile.BadZipFile):
            can_reuse = False
    if can_reuse and prior and prior.get("sha256"):
        can_reuse = hash_file(destination) == prior["sha256"]
    if can_reuse:
        result = dict(file_record)
        # A previous HTML-discovered copy may have been recorded as
        # extensionless/unsupported.  Re-open it when new MIME or signature
        # evidence makes optional extraction possible.
        prior_status = str(prior.get("extraction_status", "")) if prior else ""
        if extract_documents and prior and prior_status in {"unsupported", "not_requested", "failed", "unavailable"}:
            try:
                with destination.open("rb") as existing:
                    existing_data = existing.read(4096)
                detected_filename = _filename_with_document_extension(filename, existing_data, content_type)
                extracted_text, extraction_status = extract_document_path(
                    destination,
                    detected_filename,
                    content_type,
                    enabled=True,
                )
                reused_destination = destination
                if detected_filename != filename and destination.suffix.lower() not in {".pdf", ".docx", ".pptx"}:
                    version = metadata_version or (prior or {}).get("version_token") or _destination_version_token(
                        archive_dir,
                        source_id,
                        filename,
                        destination,
                    )
                    candidate = archive_file_path(archive_dir, source_id, detected_filename, version=version)
                    if candidate != destination:
                        candidate.parent.mkdir(parents=True, exist_ok=True)
                        os.replace(destination, candidate)
                        reused_destination = candidate
                result.update(
                    {
                        "source_id": source_id,
                        "filename": detected_filename,
                        "local_path": str(reused_destination),
                        "download_status": "reused",
                        "sha256": prior.get("sha256") or hash_file(reused_destination),
                        "extraction_status": extraction_status,
                    }
                )
                version = metadata_version or (prior or {}).get("version_token") or _destination_version_token(
                    archive_dir,
                    source_id,
                    detected_filename,
                    reused_destination,
                )
                if version:
                    result["version_token"] = str(version)
                if extracted_text is not None:
                    result["extracted_text"] = extracted_text
                # A prior failed snapshot may have carried a stale download
                # warning into this fresh successful reuse.
                result.pop("warning", None)
                return result
            except OSError:
                pass
        result.update(
            {
                "source_id": source_id,
                "filename": filename,
                "local_path": str(destination),
                "download_status": "reused",
            }
        )
        if prior and prior.get("sha256"):
            result["sha256"] = prior["sha256"]
        version = metadata_version or (prior or {}).get("version_token") or _destination_version_token(
            archive_dir,
            source_id,
            filename,
            destination,
        )
        if version:
            result["version_token"] = str(version)
        if prior and prior.get("extracted_text") is not None:
            result["extracted_text"] = prior["extracted_text"]
            result["extraction_status"] = prior.get("extraction_status", "extracted")
        result.pop("warning", None)
        return result

    declared_size = file_record.get("size")
    try:
        if max_file_size is not None and declared_size is not None and int(declared_size) > max_file_size:
            # Canvas metadata can reject a file before opening its download
            # URL.  Keep this per-file and preserve an older local version.
            return _too_large_result(file_record, source_id, filename, prior)
    except (TypeError, ValueError):
        pass

    url = file_record.get("download_url") or file_record.get("url") or file_record.get("source_url")
    if not url:
        result = dict(file_record)
        prior_path = prior.get("local_path") if prior else None
        if prior_path and Path(str(prior_path)).is_file():
            result.update(
                {
                    "source_id": source_id,
                    "filename": filename,
                    "local_path": str(prior_path),
                    "sha256": prior.get("sha256"),
                    "download_status": "preserved",
                    "warning": "file has no download URL; prior file preserved",
                }
            )
            if prior.get("version_token"):
                result["version_token"] = prior["version_token"]
            if prior.get("extracted_text") is not None:
                result["extracted_text"] = prior["extracted_text"]
                result["extraction_status"] = prior.get("extraction_status", "extracted")
        else:
            result.update({"source_id": source_id, "filename": filename, "download_status": "failed"})
            result["warning"] = "file has no download URL"
        return result

    try:
        if downloader is None and callable(getattr(client, "download_to", None)):
            return _archive_stream(client, str(url), file_record, archive_dir, filename, content_type,
                                   prior, metadata_version, max_file_size, extract_documents)
        if downloader is not None:
            data = downloader(str(url), max_file_size)
        else:
            data = _download_bytes(client, str(url), max_file_size)
        if not isinstance(data, (bytes, bytearray, memoryview)):
            raise ArchiveError("file downloader returned an unsupported value")
        data = bytes(data)
        resolved_filename = _filename_with_document_extension(filename, data, content_type)
        if resolved_filename != filename:
            filename = resolved_filename
        data_hash = _hash_bytes(data)
        current_version = metadata_version or _version_token(file_record, data)
        if metadata_version is None and prior and prior.get("sha256") == data_hash and prior.get("local_path"):
            prior_path = Path(str(prior["local_path"]))
            if prior_path.is_file():
                merged = dict(file_record)
                merged.update(
                    {
                        "source_id": source_id,
                        "filename": filename,
                        "local_path": str(prior_path),
                        "sha256": data_hash,
                        "size": len(data),
                        "download_status": "reused",
                        "version_token": current_version,
                    }
                )
                if extract_documents:
                    extracted_text, extraction_status = extract_document_text(data, filename, content_type, enabled=True)
                    merged["extraction_status"] = extraction_status
                    if extracted_text is not None:
                        merged["extracted_text"] = extracted_text
                else:
                    merged["extraction_status"] = "not_requested"
                merged.pop("warning", None)
                return merged
        if metadata_version is None or prior_destination is None:
            destination = archive_file_path(archive_dir, source_id, filename, version=current_version)
        result = archive_bytes(
            data,
            archive_dir,
            source_id,
            filename,
            content_type=content_type,
            max_file_size=max_file_size,
            extract_documents=extract_documents,
            destination=destination,
            source_marker=_source_marker(file_record),
            version=current_version,
        )
        merged = dict(file_record)
        merged.update(result)
        if result.get("download_status") in {"downloaded", "reused"}:
            merged.pop("warning", None)
        return merged
    except FileTooLargeError:
        return _too_large_result(file_record, source_id, filename, prior)
    except Exception as exc:  # network and client errors are per-file warnings
        if type(exc).__name__ == "CanvasDownloadTooLarge":
            return _too_large_result(file_record, source_id, filename, prior)
        prior_path = prior.get("local_path") if prior else None
        if prior_path and Path(str(prior_path)).is_file():
            result = dict(file_record)
            result.update(
                {
                    "source_id": source_id,
                    "filename": filename,
                    "local_path": str(prior_path),
                    "sha256": prior.get("sha256"),
                    "download_status": "preserved",
                    "warning": f"download failed; prior file preserved ({type(exc).__name__})",
                }
            )
            if prior.get("version_token"):
                result["version_token"] = prior["version_token"]
            if prior.get("extracted_text") is not None:
                result["extracted_text"] = prior["extracted_text"]
                result["extraction_status"] = prior.get("extraction_status", "extracted")
            return result
        result = dict(file_record)
        result.update(
            {
                "source_id": source_id,
                "filename": filename,
                "download_status": "failed",
                "warning": f"download failed ({type(exc).__name__})",
            }
        )
        return result


def archive_file(
    client: Any,
    file_record: Mapping[str, Any],
    archive_dir: os.PathLike[str] | str,
    **kwargs: Any,
) -> dict[str, Any]:
    """Compatibility alias with a concise name for collection callers."""

    return download_and_archive(client, file_record, archive_dir, **kwargs)


__all__ = [
    "ArchiveError",
    "DEFAULT_MAX_FILE_SIZE",
    "FileTooLargeError",
    "archive_bytes",
    "archive_file",
    "archive_file_path",
    "download_and_archive",
    "extract_document_text",
    "safe_component",
    "source_id_for_file",
]
